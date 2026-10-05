import io
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from app.models import Customer, Transaction


class ExportService:
    """Service to export Business CRM metrics, customer tables, and transaction audit logs to Excel and PDF."""

    @staticmethod
    def generate_excel_report(business) -> io.BytesIO:
        """Generates a styled Excel (.xlsx) workbook for the given business."""
        wb = openpyxl.Workbook()

        # Styles
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=14, bold=True, color="1E293B")
        data_font = Font(name="Arial", size=10)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        # -----------------------------
        # Sheet 1: Clientes
        # -----------------------------
        ws_customers = wb.active
        ws_customers.title = "Clientes Registrados"

        ws_customers["A1"] = f"Reporte de Clientes — {business.name}"
        ws_customers["A1"].font = title_font
        ws_customers["A2"] = f"Generado: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} | Programa: {business.loyalty_type.upper()}"
        ws_customers["A2"].font = Font(name="Arial", size=9, italic=True, color="64748B")

        headers_customers = [
            "ID",
            "Nombre Completo",
            "Email",
            "Teléfono",
            "Cumpleaños",
            "Fecha Registro",
            "Puntos Actuales",
            "Sellos Actuales",
        ]

        for col_num, header_title in enumerate(headers_customers, 1):
            cell = ws_customers.cell(row=4, column=col_num)
            cell.value = header_title
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        customers = business.customers.all()
        for row_num, c in enumerate(customers, 5):
            row_data = [
                c.id,
                c.full_name,
                c.email or "—",
                c.phone or "—",
                c.birthdate.strftime("%Y-%m-%d") if c.birthdate else "—",
                c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else "—",
                round(c.current_points, 1),
                c.current_stamps,
            ]
            for col_num, val in enumerate(row_data, 1):
                cell = ws_customers.cell(row=row_num, column=col_num)
                cell.value = val
                cell.font = data_font
                cell.border = thin_border
                if col_num in (1, 5, 6, 7, 8):
                    cell.alignment = Alignment(horizontal="center")

        # Auto-adjust column widths
        for col in ws_customers.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_customers.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # -----------------------------
        # Sheet 2: Transacciones (Auditoría)
        # -----------------------------
        ws_tx = wb.create_sheet(title="Historial Transacciones")

        ws_tx["A1"] = f"Auditoría de Transacciones — {business.name}"
        ws_tx["A1"].font = title_font
        ws_tx["A2"] = f"Generado: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
        ws_tx["A2"].font = Font(name="Arial", size=9, italic=True, color="64748B")

        headers_tx = [
            "ID Tx",
            "Fecha y Hora",
            "Cliente",
            "Cajero / Staff",
            "Tipo Movimiento",
            "Monto Compra ($)",
            "Puntos",
            "Sellos",
        ]

        for col_num, header_title in enumerate(headers_tx, 1):
            cell = ws_tx.cell(row=4, column=col_num)
            cell.value = header_title
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        transactions = business.transactions.order_by(Transaction.timestamp.desc()).all()
        for row_num, tx in enumerate(transactions, 5):
            customer_name = tx.customer.full_name if tx.customer else f"ID #{tx.customer_id}"
            staff_name = tx.staff.name or tx.staff.username if tx.staff else "Sistema / Auto"
            type_label = "Abono (+)" if tx.type == "earn" else "Canje (-)"

            row_data = [
                tx.id,
                tx.timestamp.strftime("%Y-%m-%d %H:%M:%S") if tx.timestamp else "—",
                customer_name,
                staff_name,
                type_label,
                f"${tx.purchase_amount:.2f}",
                round(tx.points_amount, 1),
                tx.stamps_amount,
            ]
            for col_num, val in enumerate(row_data, 1):
                cell = ws_tx.cell(row=row_num, column=col_num)
                cell.value = val
                cell.font = data_font
                cell.border = thin_border
                if col_num in (1, 2, 5, 6, 7, 8):
                    cell.alignment = Alignment(horizontal="center")

        for col in ws_tx.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_tx.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def generate_pdf_report(business) -> io.BytesIO:
        """Generates a clean PDF summary report using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=4,
        )
        subtitle_style = ParagraphStyle(
            "DocSubtitle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=14,
        )
        section_heading = ParagraphStyle(
            "SectionHeading",
            parent=styles["Heading2"],
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#0F172A"),
            spaceBefore=12,
            spaceAfter=8,
        )
        body_style = ParagraphStyle(
            "Body",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#334155"),
        )
        body_bold = ParagraphStyle(
            "BodyBold",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#0F172A"),
            fontName="Helvetica-Bold",
        )

        elements = []

        # Title & Metadata
        elements.append(Paragraph(f"Reporte Gerencial de Lealtad — {business.name}", title_style))
        category_name = business.category.name if business.category else "General"
        status_text = "Activo" if business.is_active else "Suspendido"
        meta_text = (
            f"Categoría: <b>{category_name}</b> | Tipo: <b>{business.loyalty_type.upper()}</b> | "
            f"Estado: <b>{status_text}</b> | Fecha: <b>{datetime.now().strftime('%d/%m/%Y %H:%M')}</b>"
        )
        elements.append(Paragraph(meta_text, subtitle_style))
        elements.append(Spacer(1, 8))

        # Metrics KPI Grid
        total_cust = business.customers.count()
        total_tx = business.transactions.count()
        total_vol = sum(t.purchase_amount for t in business.transactions.all())

        kpi_data = [
            [
                Paragraph("<b>Total de Clientes</b>", body_style),
                Paragraph("<b>Total de Transacciones</b>", body_style),
                Paragraph("<b>Volumen de Compras ($)</b>", body_style),
            ],
            [
                Paragraph(f"<font size=14><b>{total_cust}</b></font>", body_style),
                Paragraph(f"<font size=14><b>{total_tx}</b></font>", body_style),
                Paragraph(f"<font size=14><b>${total_vol:.2f}</b></font>", body_style),
            ],
        ]
        kpi_table = Table(kpi_data, colWidths=[180, 180, 180])
        kpi_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#E2E8F0")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 14))

        # Top Customers Table (Recent 15)
        elements.append(Paragraph("Clientes Registrados Recientes", section_heading))
        cust_rows = [[
            Paragraph("<b>Nombre</b>", body_bold),
            Paragraph("<b>Email / Contacto</b>", body_bold),
            Paragraph("<b>Cumpleaños</b>", body_bold),
            Paragraph("<b>Saldo Puntos</b>", body_bold),
            Paragraph("<b>Saldo Sellos</b>", body_bold),
        ]]

        recent_customers = business.customers.order_by(Customer.created_at.desc()).limit(15).all()
        for c in recent_customers:
            cust_rows.append([
                Paragraph(c.full_name, body_style),
                Paragraph(c.email or c.phone or "—", body_style),
                Paragraph(c.birthdate.strftime("%d/%m/%Y") if c.birthdate else "—", body_style),
                Paragraph(f"{c.current_points:.0f} pts", body_style),
                Paragraph(f"{c.current_stamps} / {business.stamps_reward_limit}", body_style),
            ])

        if len(cust_rows) > 1:
            cust_table = Table(cust_rows, colWidths=[140, 160, 80, 80, 80])
            cust_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            elements.append(cust_table)
        else:
            elements.append(Paragraph("No hay clientes registrados aún.", body_style))

        elements.append(Spacer(1, 14))

        # Recent Transactions Audit Table (Recent 15)
        elements.append(Paragraph("Historial Reciente de Transacciones (Auditoría)", section_heading))
        tx_rows = [[
            Paragraph("<b>Fecha/Hora</b>", body_bold),
            Paragraph("<b>Cliente</b>", body_bold),
            Paragraph("<b>Cajero/Staff</b>", body_bold),
            Paragraph("<b>Tipo</b>", body_bold),
            Paragraph("<b>Monto $</b>", body_bold),
        ]]

        recent_txs = business.transactions.order_by(Transaction.timestamp.desc()).limit(15).all()
        for t in recent_txs:
            c_name = t.customer.full_name if t.customer else "—"
            s_name = t.staff.name or t.staff.username if t.staff else "Auto"
            t_type = "Abono (+)" if t.type == "earn" else "Canje (-)"
            tx_rows.append([
                Paragraph(t.timestamp.strftime("%d/%m/%Y %H:%M") if t.timestamp else "—", body_style),
                Paragraph(c_name, body_style),
                Paragraph(s_name, body_style),
                Paragraph(t_type, body_style),
                Paragraph(f"${t.purchase_amount:.2f}", body_style),
            ])

        if len(tx_rows) > 1:
            tx_table = Table(tx_rows, colWidths=[110, 150, 120, 80, 80])
            tx_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            elements.append(tx_table)
        else:
            elements.append(Paragraph("No hay transacciones registradas aún.", body_style))

        doc.build(elements)
        buffer.seek(0)
        return buffer
