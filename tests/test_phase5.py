import io
import json
import unittest
import openpyxl
from datetime import date
from app import create_app
from app.extensions import db
from app.models import Business, BusinessCategory, Customer, Staff, Transaction


class Phase5TestCase(unittest.TestCase):
    """Test suite for Phase 5: Business CRUD, Categories, Excel/PDF Exports, Inactivity Suspension, and Legal Pages."""

    def setUp(self):
        self.app = create_app("testing")
        self.client = self.app.test_client()
        self.app_context = self.app.app_context()
        self.app_context.push()
        db.create_all()

        # Seed categories
        self.cat_rest = BusinessCategory(name="Restaurante", slug="restaurante")
        self.cat_bar = BusinessCategory(name="Bar", slug="bar")
        db.session.add_all([self.cat_rest, self.cat_bar])
        db.session.commit()

        # Seed sample active business
        self.business_active = Business(
            name="Taquería Don Pepe",
            slug="taqueria-don-pepe",
            category_id=self.cat_rest.id,
            primary_color="#DC2626",
            secondary_color="#FEF08A",
            loyalty_type="stamps",
            stamps_reward_limit=6,
            is_active=True,
        )
        # Seed sample inactive (suspended) business
        self.business_suspended = Business(
            name="Bar La Cueva",
            slug="bar-la-cueva",
            category_id=self.cat_bar.id,
            primary_color="#1E293B",
            secondary_color="#94A3B8",
            loyalty_type="points",
            points_per_currency=1.0,
            is_active=False,
        )
        db.session.add_all([self.business_active, self.business_suspended])
        db.session.commit()

        # Seed customer and transactions for active business
        self.customer_pepe = Customer(
            business_id=self.business_active.id,
            full_name="Mateo Morales",
            email="mateo@test.com",
            phone="+525511223344",
            birthdate=date(1992, 6, 15),
            current_stamps=3,
        )
        self.customer_cueva = Customer(
            business_id=self.business_suspended.id,
            full_name="Raul Ramos",
            email="raul@test.com",
            current_points=100.0,
        )
        db.session.add_all([self.customer_pepe, self.customer_cueva])
        db.session.commit()

        tx = Transaction(
            business_id=self.business_active.id,
            customer_id=self.customer_pepe.id,
            type="earn",
            stamps_amount=1,
            purchase_amount=120.0,
        )
        db.session.add(tx)
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()

    def test_business_crud_and_category_assignment(self):
        """Test creating a new Business through the admin API with category and geolocation."""
        payload = {
            "name": "Cervecería Artesanal",
            "slug": "cerveceria-artesanal",
            "category_id": self.cat_bar.id,
            "primary_color": "#D97706",
            "secondary_color": "#FEF3C7",
            "loyalty_type": "points",
            "points_per_currency": 1.5,
            "latitude": 19.4300,
            "longitude": -99.1500,
            "is_active": "true",
        }
        res = self.client.post("/admin/business/create", json=payload)
        self.assertEqual(res.status_code, 201)
        data = res.get_json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["business"]["category_name"], "Bar")
        self.assertTrue(data["business"]["is_active"])

    def test_toggle_business_subscription_status(self):
        """Test toggling is_active status of a business."""
        self.assertTrue(self.business_active.is_active)

        # Toggle to inactive
        res = self.client.post(
            f"/admin/business/{self.business_active.id}/toggle-status",
            headers={"Accept": "application/json"},
        )
        self.assertEqual(res.status_code, 200)
        self.assertFalse(res.get_json()["is_active"])

        # Toggle back to active
        res2 = self.client.post(
            f"/admin/business/{self.business_active.id}/toggle-status",
            headers={"Accept": "application/json"},
        )
        self.assertEqual(res2.status_code, 200)
        self.assertTrue(res2.get_json()["is_active"])

    def test_suspended_business_blocks_registration_and_scanner(self):
        """Test that inactive businesses block customer onboarding and scanner transactions."""
        # 1. Registration attempt at /join/bar-la-cueva -> 403
        reg_res = self.client.post(
            f"/join/{self.business_suspended.slug}",
            json={"full_name": "Nuevo Cliente", "email": "nuevo@test.com"},
        )
        self.assertEqual(reg_res.status_code, 403)
        self.assertIn("suspendido", reg_res.get_json()["message"])

        # 2. Scanner lookup on suspended business -> 403
        lookup_res = self.client.post(
            "/api/scanner/lookup",
            json={"token": self.customer_cueva.qr_code_token, "business_id": self.business_suspended.id},
        )
        self.assertEqual(lookup_res.status_code, 403)

        # 3. Scanner transaction on suspended business -> 403
        tx_res = self.client.post(
            "/api/scanner/process",
            json={
                "customer_id": self.customer_cueva.id,
                "business_id": self.business_suspended.id,
                "action": "earn_points",
                "amount": 50,
            },
        )
        self.assertEqual(tx_res.status_code, 403)

    def test_excel_export_endpoint(self):
        """Test GET /admin/business/<id>/export/excel returns a valid formatted .xlsx file."""
        res = self.client.get(f"/admin/business/{self.business_active.id}/export/excel")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment", res.headers.get("Content-Disposition", ""))

        # Verify Excel workbook content with openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        self.assertIn("Clientes Registrados", wb.sheetnames)
        self.assertIn("Historial Transacciones", wb.sheetnames)

        ws_c = wb["Clientes Registrados"]
        self.assertEqual(ws_c["B5"].value, "Mateo Morales")

    def test_pdf_export_endpoint(self):
        """Test GET /admin/business/<id>/export/pdf returns a valid .pdf file."""
        res = self.client.get(f"/admin/business/{self.business_active.id}/export/pdf")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.mimetype, "application/pdf")
        self.assertIn("attachment", res.headers.get("Content-Disposition", ""))
        self.assertTrue(res.data.startswith(b"%PDF"))

    def test_legal_views(self):
        """Test /privacy and /terms legal pages render correctly."""
        res_privacy = self.client.get("/privacy")
        self.assertEqual(res_privacy.status_code, 200)
        self.assertIn(b"Aviso de Privacidad", res_privacy.data)

        res_terms = self.client.get("/terms")
        self.assertEqual(res_terms.status_code, 200)
        self.assertIn(b"T\xc3\xa9rminos y Condiciones", res_terms.data)


if __name__ == "__main__":
    unittest.main()
